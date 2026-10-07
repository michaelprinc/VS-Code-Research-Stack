from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from research_stack.documents import DocumentError, extract_document, write_docx_report, write_markdown_report
from research_stack.bundle import export_bundle
from research_stack.paths import PathPolicyError, resolve_under


def fixtures(tmp_path: Path) -> tuple[Path, Path]:
    pdf = tmp_path / "czech paper č.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    # pypdf does not author text; use a DOCX for content and an empty PDF for OCR-state testing.
    with pdf.open("wb") as stream:
        writer.write(stream)
    docx = tmp_path / "notes space.docx"
    doc = Document()
    doc.add_paragraph("Výzkumná poznámka — Unicode test.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Evidence"
    table.cell(0, 1).text = "Table value"
    doc.save(docx)
    return pdf, docx


def test_docx_unicode_and_table_locators(tmp_path: Path) -> None:
    _, docx = fixtures(tmp_path)
    extracted = extract_document(docx)
    assert "Výzkumná poznámka" in extracted.text
    assert "Evidence | Table value" in extracted.text
    assert extracted.locators[0]["item"] == 1
    assert len(extracted.sha256) == 64


def test_scanned_pdf_is_explicitly_flagged(tmp_path: Path) -> None:
    pdf, _ = fixtures(tmp_path)
    extracted = extract_document(pdf)
    assert extracted.warnings and "OCR" in extracted.warnings[0]


def test_output_never_overwrites(tmp_path: Path) -> None:
    output = tmp_path / "report.md"
    output.write_text("keep", encoding="utf-8")
    with pytest.raises(DocumentError, match="OUTPUT_EXISTS"):
        write_markdown_report(output, "Title", "Body", [])
    assert output.read_text(encoding="utf-8") == "keep"


def test_workspace_path_policy_rejects_escape(tmp_path: Path) -> None:
    root = tmp_path / "workspace"
    root.mkdir()
    with pytest.raises(PathPolicyError, match="PATH_OUTSIDE_WORKSPACE"):
        resolve_under(root, "../outside.pdf", must_exist=False)


def test_manual_bundle_copies_and_indexes_without_changing_source(tmp_path: Path) -> None:
    _, docx = fixtures(tmp_path)
    digest = __import__("hashlib").sha256(docx.read_bytes()).hexdigest()
    result = export_bundle(tmp_path / "handoff", [docx])
    assert result["sourceCount"] == 1
    copied = tmp_path / "handoff" / "01-notes space.docx"
    assert copied.read_bytes() == docx.read_bytes()
    assert __import__("hashlib").sha256(docx.read_bytes()).hexdigest() == digest
    assert digest in (tmp_path / "handoff" / "source-index.md").read_text(encoding="utf-8")


def test_docx_report_contains_sources_and_never_overwrites(tmp_path: Path) -> None:
    _, source = fixtures(tmp_path)
    digest = __import__("hashlib").sha256(source.read_bytes()).hexdigest()
    report = tmp_path / "outputs" / "report.docx"
    result = write_docx_report(report, "Evidence report", "Finding one.\nFinding two.", [{"label": "Synthetic", "reference": "page 1", "sha256": digest}])
    assert result["sha256"] == __import__("hashlib").sha256(report.read_bytes()).hexdigest()
    content = "\n".join(p.text for p in Document(report).paragraphs)
    assert "Finding one." in content and "Synthetic" in content and digest in content
    with pytest.raises(DocumentError, match="OUTPUT_EXISTS"):
        write_docx_report(report, "Replacement", "", [])
