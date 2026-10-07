"""Bounded PDF/DOCX text extraction and safe research output helpers."""

from __future__ import annotations

import hashlib
import shutil
import tempfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

MAX_INPUT_BYTES = 50 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_RESULT_CHARS = 1_000_000


class DocumentError(ValueError):
    """A bounded, user-actionable document processing error."""


@dataclass
class ExtractedDocument:
    path: str
    kind: str
    sha256: str
    text: str
    locators: list[dict[str, Any]]
    warnings: list[str]

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _check_input(path: Path) -> Path:
    if path.suffix.lower() not in {".pdf", ".docx"}:
        raise DocumentError("UNSUPPORTED_FORMAT: use a text-bearing PDF or DOCX file.")
    if not path.is_file():
        raise DocumentError("SOURCE_NOT_FOUND: selected document does not exist or is not a file.")
    if path.stat().st_size > MAX_INPUT_BYTES:
        raise DocumentError(f"INPUT_TOO_LARGE: limit is {MAX_INPUT_BYTES // (1024 * 1024)} MiB.")
    return path


def extract_document(path: Path, *, start: int = 1, count: int = 50) -> ExtractedDocument:
    """Extract a page/paragraph range; locators are 1-based and stable within the source."""
    path = _check_input(path)
    if start < 1 or count < 1 or count > 100:
        raise DocumentError("INVALID_RANGE: start must be positive and count must be 1..100.")

    warnings: list[str] = []
    locators: list[dict[str, Any]] = []
    parts: list[str] = []
    if path.suffix.lower() == ".pdf":
        try:
            reader = PdfReader(str(path), strict=False)
            if reader.is_encrypted:
                raise DocumentError("PDF_ENCRYPTED: provide an unlocked copy; the source was not changed.")
            page_total = len(reader.pages)
            if page_total > MAX_PDF_PAGES:
                raise DocumentError(f"PAGE_LIMIT: PDF has {page_total} pages; limit is {MAX_PDF_PAGES}.")
            if page_total == 0:
                raise DocumentError("PDF_EMPTY: document contains no pages.")
            end = min(page_total, start + count - 1)
            if start > page_total:
                raise DocumentError(f"INVALID_RANGE: PDF has {page_total} pages.")
            for page_number in range(start, end + 1):
                text = reader.pages[page_number - 1].extract_text() or ""
                parts.append(text)
                locators.append({"page": page_number, "characters": len(text)})
            if not any(p.strip() for p in parts):
                warnings.append("NO_TEXT_EXTRACTED: this may be a scanned PDF; OCR is not included in Phase 1.")
            kind = "pdf"
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError(f"PDF_PARSE_FAILED: {type(exc).__name__}; source was not changed.") from exc
    else:
        try:
            doc = Document(str(path))
            elements: list[tuple[str, str]] = [("paragraph", p.text) for p in doc.paragraphs]
            for table_index, table in enumerate(doc.tables, start=1):
                for row_index, row in enumerate(table.rows, start=1):
                    elements.append((f"table-{table_index}-row", " | ".join(cell.text for cell in row.cells)))
            if not elements:
                raise DocumentError("DOCX_EMPTY: no paragraphs or table rows were found.")
            end = min(len(elements), start + count - 1)
            if start > len(elements):
                raise DocumentError(f"INVALID_RANGE: DOCX has {len(elements)} extractable paragraphs/rows.")
            for index in range(start, end + 1):
                kind_name, text = elements[index - 1]
                parts.append(text)
                locators.append({"item": index, "kind": kind_name, "characters": len(text)})
            kind = "docx"
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError(f"DOCX_PARSE_FAILED: {type(exc).__name__}; source was not changed.") from exc

    result = "\n\n".join(parts)
    if len(result) > MAX_RESULT_CHARS:
        raise DocumentError("RESULT_TOO_LARGE: reduce the requested range; extraction is never silently truncated.")
    return ExtractedDocument(str(path), kind, sha256_file(path), result, locators, warnings)


def write_markdown_report(output_path: Path, title: str, body: str, sources: list[dict[str, str]]) -> dict[str, str]:
    """Create a new Markdown report. Existing files are never overwritten."""
    output_path = output_path.expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        raise DocumentError("OUTPUT_EXISTS: choose a new output filename; existing files are not overwritten.")
    source_lines = [f"- [{item['label']}]({item['reference']}) — `{item['sha256']}`" for item in sources]
    content = f"# {title.strip()}\n\n{body.rstrip()}\n\n## Sources\n\n" + "\n".join(source_lines) + "\n"
    try:
        with output_path.open("x", encoding="utf-8", newline="\n") as stream:
            stream.write(content)
    except FileExistsError as exc:
        raise DocumentError("OUTPUT_EXISTS: choose a new output filename; existing files are not overwritten.") from exc
    return {"path": str(output_path), "sha256": sha256_file(output_path)}


def write_docx_report(output_path: Path, title: str, body: str, sources: list[dict[str, str]]) -> dict[str, str]:
    """Create a new DOCX report with source references; existing files are never overwritten."""
    output_path = output_path.expanduser().resolve()
    if output_path.suffix.lower() != ".docx":
        raise DocumentError("OUTPUT_FORMAT: DOCX reports need a .docx extension.")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title.strip() or "Research report", level=1)
    for paragraph in body.splitlines():
        if paragraph.strip():
            doc.add_paragraph(paragraph)
    doc.add_heading("Sources", level=2)
    for item in sources:
        doc.add_paragraph(f"{item['label']} — {item['reference']} — SHA-256 {item['sha256']}", style="List Bullet")

    with tempfile.NamedTemporaryFile(prefix=".research-stack-", suffix=".docx", dir=output_path.parent, delete=False) as temp:
        staging = Path(temp.name)
    try:
        doc.save(staging)
        try:
            with output_path.open("xb") as destination, staging.open("rb") as source:
                shutil.copyfileobj(source, destination)
        except FileExistsError as exc:
            raise DocumentError("OUTPUT_EXISTS: choose a new output filename; existing files are not overwritten.") from exc
    finally:
        staging.unlink(missing_ok=True)
    return {"path": str(output_path), "sha256": sha256_file(output_path)}
